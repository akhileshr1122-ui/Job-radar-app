"""Turn an uploaded resume (PDF / Word / text) into profile/resume.json.

With AI available (Claude API key or Claude Pro/Max token) the resume is read by Claude into the full structure.
Without AI a simple section parser fills what it can; the person finishes it in the app's profile editor.
"""
import json
import os
import re

from tailor import ai_available, call_claude

SCHEMA_HINT = """{
  "contact": {"name": "", "display_name": "NAME IN CAPS", "location": "City, Province/State, Country", "phone": "", "email": "",
              "linkedin": "", "work_authorization": ""},
  "headlines": {"main": "one-line professional headline", "alt1": "variant for a different kind of role", "alt2": "another variant"},
  "summary_base": "3-4 sentence summary",
  "summary_facts": ["each strongest, quantified achievement as a short fact"],
  "skills": {"group name": ["skill", "skill"]},
  "experience": [{"id": "shortid", "company": "", "location": "", "title": "", "start": "MM/YYYY", "end": "MM/YYYY or Present",
                  "core": ["2-3 most important bullets, always shown"],
                  "flex": [{"text": "other bullet", "tags": ["lowercase-topic", "tags"]}]}],
  "projects": [{"name": "", "text": "", "tags": []}],
  "certifications": [""],
  "education": [{"credential": "", "school": "", "dates": ""}],
  "application_answers": {"full_name": "", "email": "", "phone": "", "city": "", "linkedin": "", "years_experience": ""}
}"""

IMPORT_PROMPT = """Convert this resume into JSON with exactly this structure (newest job first):

{schema}

Rules: copy facts exactly, never invent anything; keep every number as written; put each role's 2-3 strongest bullets in "core" and the rest in "flex";
group skills into 3-6 sensible groups; leave fields empty when the resume doesn't say. Return ONLY the JSON.

RESUME TEXT:
{text}"""


def extract_text(path):
    low = path.lower()
    if low.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if low.endswith(".docx"):
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    with open(path, encoding="utf-8", errors="ignore") as fh:
        return fh.read()


def ai_import(text, cfg):
    out = call_claude("You convert resumes into structured JSON. Output JSON only.",
                      IMPORT_PROMPT.format(schema=SCHEMA_HINT, text=text[:30000]), cfg, max_tokens=8000)
    if not out:
        return None
    m = re.search(r"\{.*\}", out, re.S)
    try:
        return json.loads(m.group(0)) if m else None
    except ValueError:
        return None


SECTION_WORDS = ("summary", "profile", "personal profile", "professional summary", "about me", "objective", "experience",
                 "work experience", "professional experience", "employment", "employment history", "education", "skills",
                 "technical skills", "core skills", "key skills", "tools", "tools worked on", "certification", "certifications",
                 "licenses", "projects", "achievements", "awards")
SECTION_RX = re.compile(r"^\s*(" + "|".join(sorted(map(re.escape, SECTION_WORDS), key=len, reverse=True)) + r")\b.{0,30}$", re.I)


class _Sec:
    @staticmethod
    def match(line):
        return SECTION_RX.match(line) if len(line.strip()) < 45 else None


SECTION = _Sec()
DATES = re.compile(r"((?:\d{1,2}/)?\d{4}|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.? \d{4})\s*[-–—to]+\s*"
                   r"((?:\d{1,2}/)?\d{4}|(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.? \d{4}|present|current|now)", re.I)
BULLET = re.compile(r"^\s*[•\-\*▪●◦·]\s*")


def heuristic_import(text):
    lines = [l.rstrip() for l in text.splitlines()]
    lines = [l for l in lines if l.strip()]
    email_m = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    email = email_m.group(0) if email_m else ""
    phone_m = re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)
    link_m = re.search(r"(https?://)?(www\.)?linkedin\.com/in/[\w-]+", text, re.I)
    name = lines[0].strip() if lines else ""
    sections, cur = {"_top": []}, "_top"
    for l in lines[1:]:
        m = SECTION.match(l)
        if m:
            cur = m.group(1).lower()
            sections.setdefault(cur, [])
            continue
        sections.setdefault(cur, []).append(l)

    def sec(*names):
        for k, v in sections.items():
            if any(k.startswith(n) for n in names):
                return v
        return []

    summary = " ".join(BULLET.sub("", l) for l in sec("summary", "profile", "personal profile", "professional summary", "about", "objective"))
    skills = [BULLET.sub("", s).strip() for l in sec("skills", "technical skills", "core skills", "key skills", "tools") for s in re.split(r"[,|;]", l)]
    skills = [s for s in skills if 1 < len(s) < 60]
    certs = [BULLET.sub("", l).strip() for l in sec("certification", "licenses")]

    roles, role = [], None
    exp = [re.sub(r"\s+", " ", l).strip() for l in sec("experience", "work experience", "professional experience", "employment")]
    for i, l in enumerate(exp):
        d = DATES.search(l)
        nxt_has_date = i + 1 < len(exp) and DATES.search(exp[i + 1])
        if d:
            head = DATES.sub("", l).strip(" |,–-")
            prev = exp[i - 1] if i > 0 and not BULLET.match(exp[i - 1]) and not DATES.search(exp[i - 1]) else ""
            if role and prev and role["bullets"] and role["bullets"][-1] == BULLET.sub("", prev).strip():
                role["bullets"].pop()  # that line was this role's title, not the last role's bullet
            title, company = (prev, head) if prev else (head, "")
            role = {"title": title, "company": company, "location": "", "start": d.group(1), "end": d.group(2).title(), "bullets": []}
            roles.append(role)
        elif role and not nxt_has_date and len(l) > 25:
            role["bullets"].append(BULLET.sub("", l).strip())
        elif role and not nxt_has_date and role["bullets"] and len(l) <= 25 and not BULLET.match(l):
            role["bullets"][-1] += " " + l
    experience = []
    for n, r in enumerate(roles):
        r["company"] = re.sub(r"^company\s*[-:]\s*", "", r["company"], flags=re.I)
        rid = re.sub(r"[^a-z0-9]", "", (r["company"] or r["title"]).lower())[:16] or f"role{n}"
        experience.append({"id": rid, "company": r["company"], "location": "", "title": r["title"], "start": r["start"], "end": r["end"],
                           "core": r["bullets"][:2], "flex": [{"text": b, "tags": []} for b in r["bullets"][2:]]})
    first_title = experience[0]["title"] if experience else ""
    location = ""
    for l in lines[1:6]:
        m = re.search(r"([A-Z][A-Za-z .'-]+),\s*([A-Z]\d[A-Z]\s?\d[A-Z]\d)", l)  # Canadian postal code
        if m:
            location = f"{m.group(1).split(',')[-1].strip()}, Canada"
            break
        m = re.search(r"([A-Z][A-Za-z .'-]+),\s*(ON|QC|BC|AB|MB|SK|NS|NB|NL|PE|[A-Z]{2})\b", l)
        if m and not re.search(r"@|\d{3}[- ]\d{3}", m.group(0)):
            location = f"{m.group(1).strip()}, {m.group(2)}"
            break
    education = []
    for l in [re.sub(r"\s+", " ", x).strip() for x in sec("education")]:
        d = DATES.search(l)
        if d and education and not education[-1]["dates"]:
            education[-1]["school"] = DATES.sub("", l).strip(" |,–-")
            education[-1]["dates"] = f"{d.group(1)} – {d.group(2).title()}"
        elif not re.match(r"(?i)relevant|courses|modules|gpa", l) and len(education) < 6:
            education.append({"credential": BULLET.sub("", l), "school": "", "dates": ""})
    return {
        "contact": {"name": name.title() if name.isupper() else name, "display_name": name.upper(), "location": location,
                    "phone": phone_m.group(1).strip() if phone_m else "", "email": email or "",
                    "linkedin": ("https://" + link_m.group(0).split("://")[-1]) if link_m else "", "work_authorization": ""},
        "headlines": {"main": first_title},
        "summary_base": summary[:900],
        "summary_facts": [],
        "skills": {"Skills": skills[:40]} if skills else {"Skills": []},
        "experience": experience,
        "projects": [],
        "certifications": certs,
        "education": education,
        "application_answers": {"full_name": name.title() if name.isupper() else name, "email": email or "",
                                "phone": phone_m.group(1).strip() if phone_m else ""},
    }


def normalize(r):
    """Fill defaults so the rest of the engine can rely on every key."""
    r.setdefault("contact", {})
    for k in ("name", "display_name", "location", "phone", "email", "linkedin", "work_authorization"):
        r["contact"].setdefault(k, "")
    if not r["contact"]["display_name"]:
        r["contact"]["display_name"] = r["contact"]["name"].upper()
    r.setdefault("headlines", {})
    if not any(r["headlines"].values()):
        r["headlines"] = {"main": (r.get("experience") or [{}])[0].get("title", "")}
    for k, d in (("summary_base", ""), ("summary_facts", []), ("skills", {}), ("experience", []), ("projects", []),
                 ("certifications", []), ("education", []), ("application_answers", {})):
        r.setdefault(k, d)
    seen = set()
    for i, e in enumerate(r["experience"]):
        rid = e.get("id") or re.sub(r"[^a-z0-9]", "", (e.get("company") or "role").lower())[:16] or f"role{i}"
        while rid in seen:
            rid += str(i)
        seen.add(rid)
        e["id"] = rid
        for k in ("company", "location", "title", "start", "end"):
            e.setdefault(k, "")
        e.setdefault("core", [])
        e["flex"] = [f if isinstance(f, dict) else {"text": f, "tags": []} for f in e.get("flex", [])]
        for f in e["flex"]:
            f.setdefault("tags", [])
    for p in r["projects"]:
        p.setdefault("tags", [])
    return r


def import_resume(path, cfg):
    text = extract_text(path)
    if len(text.strip()) < 100:
        raise ValueError("Couldn't read text from that file (is it a scanned image?)")
    data = ai_import(text, cfg) if ai_available() else None
    method = "ai" if data else "basic"
    data = normalize(data or heuristic_import(text))
    data["_readme"] = ("Master resume data. Core bullets, companies, titles, dates, education and certifications are never changed "
                       "by tailoring; only headlines, summary, skill order and optional (flex) bullets are chosen per job.")
    data["_imported"] = {"method": method, "file": os.path.basename(path)}
    return data


def setup_search(cfg, resume):
    """First-time search settings from a freshly imported resume."""
    clean = lambda t: re.sub(r"\s*\(.*?\)|/.*$", "", t).strip()
    titles = list(dict.fromkeys(clean(e["title"]) for e in resume["experience"] if e.get("title")))[:4]
    if titles:
        cfg["queries"] = list(dict.fromkeys(titles + [clean(h) for h in resume.get("headlines", {}).values() if h and len(h) < 60][:2]))
    loc = resume.get("contact", {}).get("location", "")
    if loc:
        cfg.setdefault("locations", {})["search_locations"] = [{"place": loc, "radius_km": 50}]
        from util import classify_location
        where = classify_location(loc)[0]
        if where in ("AU", "IN"):
            cfg["locations"]["country"] = where
        elif re.search(r"\b(usa|united states|us)\b", loc, re.I) or re.search(r",\s*[A-Z]{2}\s*$", loc) and not re.search(
                r",\s*(ON|QC|BC|AB|MB|SK|NS|NB|NL|PE)\s*$", loc):
            cfg["locations"]["country"] = "US"
    cfg["title_terms"]["required"] = []
    cfg["description_keywords"] = {}
    cfg.pop("_setup", None)
    return cfg

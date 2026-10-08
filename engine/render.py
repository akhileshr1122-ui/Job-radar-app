"""Render a tailored resume to PDF (reportlab) and DOCX (python-docx). Single column, ATS-friendly."""
import re

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (HRFlowable, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

ACCENT = "#0F4C5C"


def _sections(t, resume):
    c = resume["contact"]
    roles = {e["id"]: e for e in resume["experience"]}
    exp = []
    for e in t["experience"]:
        r = roles[e["id"]]
        where = r.get("company", "") + (f", {r['location']}" if r.get("location") else "")
        exp.append({"title": r["title"], "where": where, "dates": f"{r['start']} – {r['end']}", "bullets": e["bullets"]})
    projects = [p for p in resume["projects"] if p["name"] in t.get("projects", [])]
    contact_line = " · ".join(x for x in [c.get("location"), c.get("phone"), c.get("email"), (c.get("linkedin") or "").replace("https://www.", "")] if x)
    return c, exp, projects, contact_line


def _esc(s):
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_pdf(t, resume, path):
    c, exp, projects, contact_line = _sections(t, resume)
    st = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=20, leading=23, textColor=colors.HexColor("#111111")),
        "head": ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor(ACCENT), spaceBefore=2),
        "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=8.8, leading=11, textColor=colors.HexColor("#444444")),
        "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=colors.HexColor(ACCENT), spaceBefore=9, spaceAfter=2),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=9.3, leading=12.2, alignment=TA_LEFT),
        "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=9.8, leading=12),
        "dates": ParagraphStyle("dates", fontName="Helvetica", fontSize=9, leading=12, alignment=2, textColor=colors.HexColor("#444444")),
        "where": ParagraphStyle("where", fontName="Helvetica-Oblique", fontSize=9.1, leading=11.5, textColor=colors.HexColor("#333333")),
        "bullet": ParagraphStyle("bullet", fontName="Helvetica", fontSize=9.2, leading=11.8),
    }
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.6 * inch, rightMargin=0.6 * inch, topMargin=0.5 * inch,
                            bottomMargin=0.5 * inch, title=f"{c['name']} – Resume", author=c["name"])
    W = LETTER[0] - 1.2 * inch
    f = [Paragraph(_esc(c.get("display_name") or c.get("name", "")), st["name"])]
    if t.get("headline"):
        f.append(Paragraph(_esc(t["headline"]), st["head"]))
    for line in (contact_line, c.get("work_authorization", "")):
        if line:
            f.append(Paragraph(_esc(line), st["contact"]))

    def section(title):
        f.append(Paragraph(title.upper(), st["h"]))
        f.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor(ACCENT), spaceBefore=0, spaceAfter=4))

    if t.get("summary"):
        section("Summary")
        f.append(Paragraph(_esc(t["summary"]), st["body"]))
    if t.get("skills"):
        section("Core skills")
        f.append(Paragraph(_esc("  •  ".join(t["skills"])), st["body"]))
    section("Professional experience")
    for e in exp:
        head = Table([[Paragraph(_esc(e["title"]), st["role"]), Paragraph(_esc(e["dates"]), st["dates"])]],
                     colWidths=[W * 0.72, W * 0.28])
        head.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                  ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                                  ("VALIGN", (0, 0), (-1, -1), "BOTTOM")]))
        bullets = ListFlowable([ListItem(Paragraph(_esc(b), st["bullet"]), leftIndent=10, value="•") for b in e["bullets"]],
                               bulletType="bullet", start="•", leftIndent=10, bulletFontSize=8)
        f.append(KeepTogether([head, Paragraph(_esc(e["where"]), st["where"]), Spacer(1, 2)]))
        f.append(bullets)
    if projects:
        section("Selected projects (built with AI-assisted development)")
        for p in projects:
            f.append(Paragraph(f"<b>{_esc(p['name'])}</b> – {_esc(p['text'])}", st["bullet"]))
            f.append(Spacer(1, 2))
    if resume.get("certifications"):
        section("Certifications")
        f.append(Paragraph(_esc("  •  ".join(resume["certifications"])), st["body"]))
    if resume.get("education"):
        section("Education")
    for ed in resume.get("education", []):
        rest = ", ".join(x for x in [ed.get("school"), ed.get("dates")] if x)
        f.append(Paragraph(f"<b>{_esc(ed['credential'])}</b>" + (f" – {_esc(rest)}" if rest else ""), st["body"]))
    doc.build(f)


def _hr(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": ACCENT.lstrip("#")}.items():
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)


def to_docx(t, resume, path):
    c, exp, projects, contact_line = _sections(t, resume)
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(0.6)
        s.top_margin = s.bottom_margin = Inches(0.5)
    base = d.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)
    base.paragraph_format.space_after = Pt(0)
    accent = RGBColor.from_string(ACCENT.lstrip("#"))

    p = d.add_paragraph()
    r = p.add_run(c["display_name"]); r.bold = True; r.font.size = Pt(20)
    if t.get("headline"):
        p = d.add_paragraph(); r = p.add_run(t["headline"]); r.bold = True; r.font.color.rgb = accent; r.font.size = Pt(11)
    for line in (contact_line, c.get("work_authorization", "")):
        if line:
            d.add_paragraph(line).runs[0].font.size = Pt(9)

    def section(title):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title.upper()); r.bold = True; r.font.color.rgb = accent; r.font.size = Pt(10.5)
        _hr(p)

    if t.get("summary"):
        section("Summary"); d.add_paragraph(t["summary"])
    if t.get("skills"):
        section("Core skills"); d.add_paragraph("  •  ".join(t["skills"]))
    section("Professional experience")
    for e in exp:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(e["title"]); r.bold = True
        p.add_run("\t" + e["dates"]).font.size = Pt(9.5)
        p = d.add_paragraph(); r = p.add_run(e["where"]); r.italic = True
        for b in e["bullets"]:
            d.add_paragraph(b, style="List Bullet")
    if projects:
        section("Selected projects (built with AI-assisted development)")
        for pr in projects:
            p = d.add_paragraph(style="List Bullet"); r = p.add_run(pr["name"] + " – "); r.bold = True; p.add_run(pr["text"])
    if resume.get("certifications"):
        section("Certifications"); d.add_paragraph("  •  ".join(resume["certifications"]))
    if resume.get("education"):
        section("Education")
    for ed in resume.get("education", []):
        p = d.add_paragraph(); r = p.add_run(ed["credential"]); r.bold = True
        rest = ", ".join(x for x in [ed.get("school"), ed.get("dates")] if x)
        if rest:
            p.add_run(f" – {rest}")
    d.core_properties.author = c["name"]
    d.core_properties.title = f"{c['name']} – Resume"
    d.save(path)


def file_stem(resume, job=None):
    name = re.sub(r"[^A-Za-z0-9]+", "_", resume.get("contact", {}).get("name", "Resume")).strip("_")
    if not job:
        return f"{name}_Resume"
    co = re.sub(r"[^A-Za-z0-9]+", "", job["company"].title())[:24] or "Company"
    return f"{name}_{co}"
